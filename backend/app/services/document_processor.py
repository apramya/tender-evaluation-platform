"""
Document processing service for text extraction and OCR
"""
import logging
import os
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import asyncio
from PIL import Image
import pytesseract
import PyPDF2
import pdfplumber
from docx import Document as DocxDocument
try:
    import fitz  # PyMuPDF
except ImportError:  # pragma: no cover - optional dependency in older installs
    fitz = None

from app.utils.config import settings

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Service for processing various document types"""

    MIN_TEXT_CHARS_FOR_PDF = 120
    
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Extract text from PDF and OCR scanned pages when text extraction is weak."""
        text_content = []
        
        try:
            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages):
                    text = page.extract_text()
                    if text:
                        text_content.append(f"[PAGE {page_num + 1}]\n{text}\n")
                    
                    # Extract tables if any
                    tables = page.extract_tables()
                    if tables:
                        for table in tables:
                            text_content.append("\n[TABLE]\n")
                            for row in table:
                                text_content.append(" | ".join(str(cell) if cell else "" for cell in row) + "\n")
        
        except Exception as e:
            logger.error(f"Error extracting text from PDF {file_path}: {e}")
            # Fallback to PyPDF2
            try:
                with open(file_path, 'rb') as pdf_file:
                    pdf_reader = PyPDF2.PdfReader(pdf_file)
                    for page in pdf_reader.pages:
                        text = page.extract_text()
                        text_content.append(text)
            except Exception as e2:
                logger.error(f"PyPDF2 also failed: {e2}")
                return ""
        
        extracted = "\n".join(text_content).strip()
        if DocumentProcessor._is_weak_pdf_text(extracted):
            ocr_text = DocumentProcessor.extract_text_from_scanned_pdf(file_path)
            if ocr_text:
                if extracted:
                    return f"{extracted}\n\n[OCR FALLBACK]\n{ocr_text}"
                return ocr_text

        return extracted

    @staticmethod
    def _is_weak_pdf_text(text: str) -> bool:
        words = [word for word in text.split() if any(ch.isalnum() for ch in word)]
        return len(text.strip()) < DocumentProcessor.MIN_TEXT_CHARS_FOR_PDF or len(words) < 25

    @staticmethod
    def extract_text_from_scanned_pdf(file_path: str) -> str:
        """Render PDF pages to images and run OCR for scanned/image-only PDFs."""
        if not settings.USE_TESSERACT:
            logger.warning("Tesseract OCR is disabled; scanned PDF OCR skipped")
            return ""
        if fitz is None:
            logger.warning("PyMuPDF is not installed; scanned PDF OCR skipped")
            return ""

        text_content = []
        try:
            with fitz.open(file_path) as pdf:
                for page_num, page in enumerate(pdf):
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    if image.mode != "L":
                        image = image.convert("L")
                    text = pytesseract.image_to_string(image)
                    if text.strip():
                        text_content.append(f"[PAGE {page_num + 1} OCR]\n{text.strip()}\n")
        except Exception as e:
            logger.error(f"Error OCRing scanned PDF {file_path}: {e}")

        return "\n".join(text_content).strip()
    
    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from DOCX file"""
        text_content = []
        
        try:
            doc = DocxDocument(file_path)
            
            for para in doc.paragraphs:
                if para.text.strip():
                    text_content.append(para.text)
            
            # Extract tables
            for table in doc.tables:
                text_content.append("\n[TABLE]\n")
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    text_content.append(row_text)
        
        except Exception as e:
            logger.error(f"Error extracting text from DOCX {file_path}: {e}")
            return ""
        
        return "\n".join(text_content)
    
    @staticmethod
    def extract_text_from_image(file_path: str) -> str:
        """Extract text from image using OCR"""
        text = ""
        
        try:
            if not settings.USE_TESSERACT:
                logger.warning("Tesseract OCR is disabled")
                return text
            
            image = Image.open(file_path)
            
            # Preprocess image for better OCR
            # Convert to grayscale
            if image.mode != 'L':
                image = image.convert('L')
            
            # Extract text using Tesseract
            text = pytesseract.image_to_string(image)
            
        except Exception as e:
            logger.error(f"Error extracting text from image {file_path}: {e}")
        
        return text
    
    @staticmethod
    def extract_text_from_text_file(file_path: str) -> str:
        """Extract text from plain text file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except UnicodeDecodeError:
            # Try with different encoding
            try:
                with open(file_path, 'r', encoding='latin-1') as f:
                    return f.read()
            except Exception as e:
                logger.error(f"Error reading text file {file_path}: {e}")
                return ""
    
    @staticmethod
    async def process_document(file_path: str, file_type: str) -> Tuple[str, bool]:
        """
        Process document and extract text based on file type
        Returns: (extracted_text, success)
        """
        if not os.path.exists(file_path):
            logger.error(f"File not found: {file_path}")
            return "", False
        
        try:
            if file_type.lower() == 'pdf':
                text = DocumentProcessor.extract_text_from_pdf(file_path)
            elif file_type.lower() in ['docx', 'doc']:
                text = DocumentProcessor.extract_text_from_docx(file_path)
            elif file_type.lower() in ['png', 'jpg', 'jpeg', 'tiff']:
                text = DocumentProcessor.extract_text_from_image(file_path)
            elif file_type.lower() in ['txt']:
                text = DocumentProcessor.extract_text_from_text_file(file_path)
            else:
                logger.warning(f"Unsupported file type: {file_type}")
                return "", False
            
            if text:
                return text, True
            else:
                logger.warning(f"No text extracted from {file_path}")
                return "", False
        
        except Exception as e:
            logger.error(f"Error processing document {file_path}: {e}")
            return "", False
    
    @staticmethod
    def chunk_text(
        text: str,
        chunk_size: int = settings.CHUNK_SIZE,
        overlap: int = settings.CHUNK_OVERLAP,
    ) -> List[str]:
        """Split text into overlapping chunks"""
        chunks = []
        
        if len(text) <= chunk_size:
            return [text]
        
        for i in range(0, len(text), chunk_size - overlap):
            chunk = text[i:i + chunk_size]
            chunks.append(chunk)
        
        return chunks
    
    @staticmethod
    def extract_metadata_from_text(text: str) -> Dict:
        """Extract basic metadata from text"""
        metadata = {
            "word_count": len(text.split()),
            "character_count": len(text),
            "line_count": len(text.split('\n')),
            "has_tables": "[TABLE]" in text,
            "has_multiple_pages": "[PAGE" in text,
        }
        return metadata
